
"""
Convert cleaned widened coral data into a human-friendly Excel workbook.

This script:
1. Loads the cleaned widened CSV
2. Renames columns for human readability/backwards compatibility
3. Creates an ALL_SITES worksheet
4. Creates individual worksheets for each site/habitat combination
5. Applies conditional formatting to dynamics and special values
6. Saves the final human-friendly Excel workbook
"""

# =============================================================================
# CONFIGURATION
# =============================================================================

# The MOST RECENT survey year.
# Change this value when processing a new survey year.
SURVEY_YEAR = 2025

# Years represented in the widened dataset.
START_YEAR = 2013

# Input/output directory.
DATA_OUTPUTS_DIR = "../data_outputs"

# Excel sheet containing the complete dataset.
ALL_SITES_SHEET = "ALL_SITES"


# =============================================================================
# IMPORTS
# =============================================================================

from pathlib import Path
import re

import pandas as pd

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.formatting.rule import (
    CellIsRule,
    FormulaRule,
    Rule,
)


# =============================================================================
# FILE PATHS
# =============================================================================

data_dir = Path(DATA_OUTPUTS_DIR)

input_file = (
    data_dir
    / f"coral_clean_wide_{START_YEAR}-{SURVEY_YEAR}.csv"
)

output_file = (
    data_dir
    / f"coral_clean_wide_human-friendly_{START_YEAR}-{SURVEY_YEAR}.xlsx"
)


# =============================================================================
# LOAD DATA
# =============================================================================

coral_clean_wide = pd.read_csv(input_file)


# =============================================================================
# COLUMN NAME CHANGES
# =============================================================================

# Rename:
#   height_XXXX       -> HXXXX
#   width_XXXX        -> WXXXX
#   length_XXXX       -> LXXXX
#   observer_XXXX     -> ObsXXXX
#   note_extra_XXXX   -> note_XXXX_extra

def rename_columns(columns):
    """Apply the human-friendly column naming conventions."""

    renamed = []

    for column in columns:

        if column.startswith("height_"):
            column = column.replace("height_", "H", 1)

        elif column.startswith("width_"):
            column = column.replace("width_", "W", 1)

        elif column.startswith("length_"):
            column = column.replace("length_", "L", 1)

        elif column.startswith("observer_"):
            column = column.replace("observer_", "Obs", 1)

        # note_extra_XXXX -> note_XXXX_extra
        column = re.sub(
            r"^note_extra_(\d+)$",
            r"note_\1_extra",
            column
        )

        renamed.append(column)

    return renamed


coral_clean_wide.columns = rename_columns(coral_clean_wide.columns)


# =============================================================================
# CREATE WORKBOOK
# =============================================================================

wb = Workbook()

# Remove the default worksheet.
default_sheet = wb.active
wb.remove(default_sheet)


# =============================================================================
# HELPER FUNCTIONS
# =============================================================================

def write_dataframe_to_sheet(ws, dataframe):
    """
    Write a pandas DataFrame into an openpyxl worksheet.

    Header is written on row 1 and data begins on row 2.
    """

    # Header
    for col_index, column_name in enumerate(dataframe.columns, start=1):
        ws.cell(
            row=1,
            column=col_index,
            value=column_name
        )

    # Data
    for row_index, row in enumerate(
        dataframe.itertuples(index=False, name=None),
        start=2
    ):
        for col_index, value in enumerate(row, start=1):

            # Convert pandas NaN to None so Excel gets an empty cell.
            if pd.isna(value):
                value = None

            ws.cell(
                row=row_index,
                column=col_index,
                value=value
            )


def make_sheet_name(site, habitat):
    """
    Create the sheet name used by the original R script:

        paste0(site, habitat)

    Excel sheet names have a maximum length of 31 characters and cannot
    contain certain characters, so clean the resulting name.
    """

    sheet_name = f"{site}{habitat}"

    # Characters prohibited in Excel sheet names.
    sheet_name = re.sub(r'[\[\]:*?/\\]', "_", sheet_name)

    # Excel's maximum worksheet name length.
    sheet_name = sheet_name[:31]

    return sheet_name


def excel_column_letter(column_number):
    """Convert a 1-based Excel column number to its letter."""

    letters = ""

    while column_number:
        column_number, remainder = divmod(column_number - 1, 26)
        letters = chr(65 + remainder) + letters

    return letters


# =============================================================================
# CREATE ALL_SITES SHEET
# =============================================================================

ws = wb.create_sheet(ALL_SITES_SHEET)

write_dataframe_to_sheet(
    ws,
    coral_clean_wide
)


# =============================================================================
# CREATE SITE/HABITAT SHEETS
# =============================================================================

site_habitat = (
    coral_clean_wide[["site", "habitat"]]
    .drop_duplicates()
)

for _, combination in site_habitat.iterrows():

    site = combination["site"]
    habitat = combination["habitat"]

    sheet_name = make_sheet_name(site, habitat)

    # Filter the data for this site/habitat.
    subset = coral_clean_wide[
        (coral_clean_wide["site"] == site)
        & (coral_clean_wide["habitat"] == habitat)
    ]

    ws = wb.create_sheet(sheet_name)

    write_dataframe_to_sheet(
        ws,
        subset
    )


# =============================================================================
# CONDITIONAL FORMATTING STYLES
# =============================================================================

# Current values described in common-speak
# Recruitment: light blue
# Growth: light green
# Shrinkage: dark green
# Death: red
# Fission: light blue
# Fusion: light bluish purple
# Multiple: black
# Missing: brown

styles = {

    "Recruitment": {
        "font": Font(color="000000"),
        "fill": PatternFill(
            fill_type="solid",
            fgColor="159FEC"
        ),
    },

    "Growth": {
        "font": Font(color="000000"),
        "fill": PatternFill(
            fill_type="solid",
            fgColor="17A43F"
        ),
    },

    "Shrinkage": {
        "font": Font(color="FFFFFF"),
        "fill": PatternFill(
            fill_type="solid",
            fgColor="006400"  # darkgreen
        ),
    },

    "Death": {
        "font": Font(color="000000"),
        "fill": PatternFill(
            fill_type="solid",
            fgColor="FB0006"
        ),
    },

    "Fission": {
        "font": Font(color="000000"),
        "fill": PatternFill(
            fill_type="solid",
            fgColor="A1BCD6"
        ),
    },

    "Fusion": {
        "font": Font(color="000000"),
        "fill": PatternFill(
            fill_type="solid",
            fgColor="B07EDD"
        ),
    },

    "MultipleDynamics": {
        "font": Font(color="FFFFFF"),
        "fill": PatternFill(
            fill_type="solid",
            fgColor="000000"
        ),
    },

    "NaValue": {
        "font": Font(color="000000"),
        "fill": PatternFill(
            fill_type="solid",
            fgColor="FFFF00"
        ),
    },

    "UkValue": {
        "font": Font(color="000000"),
        "fill": PatternFill(
            fill_type="solid",
            fgColor="FFA500"
        ),
    },

    "MissingData": {
        "font": Font(color="FFFFFF"),
        "fill": PatternFill(
            fill_type="solid",
            fgColor="A52A2A"  # brown
        ),
    },
}


# =============================================================================
# APPLY CONDITIONAL FORMATTING
# =============================================================================

# Find all dynamics columns.
#
# Equivalent to:
#
#   which(grepl("^dynamics", colnames(coral_clean_wide)))
#
dynamics_columns = [
    index + 1
    for index, column in enumerate(coral_clean_wide.columns)
    if column.startswith("dynamics")
]


from openpyxl.formatting.rule import CellIsRule, FormulaRule


# TO-DO Address ISSUE: It appears that conditional formatting isn't working.
# Colors are not appearing
def apply_conditional_formatting(ws):
    """
    Apply conditional formatting to a worksheet.
    """

    first_data_row = 2
    last_data_row = ws.max_row

    if last_data_row < first_data_row:
        return

    # ========================================================================
    # DYNAMICS COLUMNS
    # ========================================================================

    for column_number in dynamics_columns:

        column_letter = excel_column_letter(column_number)

        cell_range = (
            f"{column_letter}{first_data_row}:"
            f"{column_letter}{last_data_row}"
        )

        # --------------------------------------------------------------------
        # Exact matches
        # --------------------------------------------------------------------

        ws.conditional_formatting.add(
            cell_range,
            CellIsRule(
                operator="equal",
                formula=['"Recruitment"'],
                font=styles["Recruitment"]["font"],
                fill=styles["Recruitment"]["fill"],
            )
        )

        ws.conditional_formatting.add(
            cell_range,
            CellIsRule(
                operator="equal",
                formula=['"Growth"'],
                font=styles["Growth"]["font"],
                fill=styles["Growth"]["fill"],
            )
        )

        ws.conditional_formatting.add(
            cell_range,
            CellIsRule(
                operator="equal",
                formula=['"Shrinkage"'],
                font=styles["Shrinkage"]["font"],
                fill=styles["Shrinkage"]["fill"],
            )
        )

        ws.conditional_formatting.add(
            cell_range,
            CellIsRule(
                operator="equal",
                formula=['"Death"'],
                font=styles["Death"]["font"],
                fill=styles["Death"]["fill"],
            )
        )

        ws.conditional_formatting.add(
            cell_range,
            CellIsRule(
                operator="equal",
                formula=['"Missing Data"'],
                font=styles["MissingData"]["font"],
                fill=styles["MissingData"]["fill"],
            )
        )

        # --------------------------------------------------------------------
        # Fission
        # Begins with "Fission"
        # --------------------------------------------------------------------

        ws.conditional_formatting.add(
            cell_range,
            FormulaRule(
                formula=[
                    f'LEFT({column_letter}{first_data_row},7)="Fission"'
                ],
                font=styles["Fission"]["font"],
                fill=styles["Fission"]["fill"],
            )
        )

        # --------------------------------------------------------------------
        # Fusion
        # Begins with "Fusion"
        # --------------------------------------------------------------------

        ws.conditional_formatting.add(
            cell_range,
            FormulaRule(
                formula=[
                    f'LEFT({column_letter}{first_data_row},6)="Fusion"'
                ],
                font=styles["Fusion"]["font"],
                fill=styles["Fusion"]["fill"],
            )
        )

        # --------------------------------------------------------------------
        # Multiple dynamics
        # Contains ","
        # --------------------------------------------------------------------

        ws.conditional_formatting.add(
            cell_range,
            FormulaRule(
                formula=[
                    f'ISNUMBER(SEARCH(",",{column_letter}{first_data_row}))'
                ],
                font=styles["MultipleDynamics"]["font"],
                fill=styles["MultipleDynamics"]["fill"],
            )
        )

    # ========================================================================
    # Na / UK VALUES
    # ========================================================================

    if ws.max_column >= 2:

        first_column_letter = excel_column_letter(2)
        last_column_letter = excel_column_letter(ws.max_column)

        all_data_range = (
            f"{first_column_letter}{first_data_row}:"
            f"{last_column_letter}{last_data_row}"
        )

        # --------------------------------------------------------------------
        # Na
        # --------------------------------------------------------------------

        ws.conditional_formatting.add(
            all_data_range,
            FormulaRule(
                formula=[
                    f'{first_column_letter}{first_data_row}="Na"'
                ],
                font=styles["NaValue"]["font"],
                fill=styles["NaValue"]["fill"],
            )
        )

        # --------------------------------------------------------------------
        # UK
        # --------------------------------------------------------------------

        ws.conditional_formatting.add(
            all_data_range,
            FormulaRule(
                formula=[
                    f'{first_column_letter}{first_data_row}="UK"'
                ],
                font=styles["UkValue"]["font"],
                fill=styles["UkValue"]["fill"],
            )
        )

# Apply formatting to every worksheet.
for ws in wb.worksheets:
    apply_conditional_formatting(ws)


# =============================================================================
# SAVE WORKBOOK
# =============================================================================

wb.save(output_file)

print(f"Created: {output_file}")
