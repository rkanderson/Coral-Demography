from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


# =============================================================================
# CONFIGURATION
# =============================================================================

CURRENT_SURVEY_YEAR = 2026

# The data from the previous survey year will be included alongside the
# blank fields for the current survey.
PREVIOUS_SURVEY_YEAR = CURRENT_SURVEY_YEAR - 1
TWO_YEARS_AGO = CURRENT_SURVEY_YEAR - 2

# Directory containing input/output data
DATA_DIR = Path("../data_outputs")

# Input files
TIDY_FILE = DATA_DIR / f"coral_tidy_dyn_2013-{PREVIOUS_SURVEY_YEAR}.csv"
WIDE_FILE = DATA_DIR / f"coral_clean_wide_2013-{PREVIOUS_SURVEY_YEAR}.csv"

# Output file
OUTPUT_FILE = (
    DATA_DIR / f"coral_data_entry_datasheet_{CURRENT_SURVEY_YEAR}.xlsx"
)

# Whether to exclude plots whose transect name starts with "P"
# (Currently relevant because monitoring is being done only on transects.)
EXCLUDE_PLOTS = True

# Whether to include data from 2-years-ago. This can be useful in cases where
# when an area was skipped for a year, but the corals are not connfirmed dead.
INCLUDE_2_YEARS_AGO = True


# =============================================================================
# LOAD DATA
# =============================================================================

# Load tidy dataset containing dynamics
coral_df_tidy = pd.read_csv(TIDY_FILE)

# BAND-AID FIX:
# Some transects have been entered with the letter O instead of zero.
coral_df_tidy["transect"] = coral_df_tidy["transect"].str.replace(
    "O", "0", regex=False
)


# =============================================================================
# IDENTIFY CORALS THAT HAVE EVER BEEN CONFIRMED DEAD
# =============================================================================

# Group by coral_number and determine whether any observation for that coral
# has dyn_death == 1.
coral_death_summary = (
    coral_df_tidy
    .groupby("coral_number", as_index=False)["dyn_death"]
    .apply(lambda x: (x == 1).any())
    .rename(columns={"dyn_death": "any_death"})
)


# =============================================================================
# LOAD WIDE DATASET
# =============================================================================

coral_df_wide = pd.read_csv(WIDE_FILE)

# BAND-AID FIX again:
# Correct transects containing O instead of 0.
coral_df_wide["transect"] = coral_df_wide["transect"].str.replace(
    "O", "0", regex=False
)


# =============================================================================
# REMOVE CORALS THAT HAVE BEEN CONFIRMED DEAD
# =============================================================================

coral_df_wide_no_dead = (
    coral_df_wide
    .merge(
        coral_death_summary,
        on="coral_number",
        how="left"
    )
)

# Keep only corals where any_death == False.
#
# The fillna(False) is defensive: if a coral somehow exists in the wide
# dataset but not in the death summary, we treat it as not confirmed dead.
coral_df_wide_no_dead = coral_df_wide_no_dead[
    coral_df_wide_no_dead["any_death"].fillna(False) == False
].copy()


# =============================================================================
# SELECT COLUMNS FOR DATA ENTRY SHEET
# =============================================================================

# Column names in the wide dataset
length_col = f"length_{PREVIOUS_SURVEY_YEAR}"
width_col = f"width_{PREVIOUS_SURVEY_YEAR}"
height_col = f"height_{PREVIOUS_SURVEY_YEAR}"
notes_col = f"note_{PREVIOUS_SURVEY_YEAR}"
length_col_2_yrs_ago = f"length_{TWO_YEARS_AGO}"
width_col_2_yrs_ago = f"width_{TWO_YEARS_AGO}"
height_col_2_yrs_ago = f"height_{TWO_YEARS_AGO}"
notes_col_2_yrs_ago = f"note_{TWO_YEARS_AGO}"

selected_columns = [
    "site",
    "habitat",
    "transect",
    "taxa",
    "x",
    "y",
    "z",
    length_col_2_yrs_ago,
    width_col_2_yrs_ago,
    height_col_2_yrs_ago,
    notes_col_2_yrs_ago,
    length_col,
    width_col,
    height_col,
    notes_col,
]


# If we're not including 2-years-ago, then remove those columns from the selection.
if not INCLUDE_2_YEARS_AGO:
    selected_columns = [
        col for col in selected_columns if not col.endswith(f"_{TWO_YEARS_AGO}")
    ]
  

coral_df_selected = coral_df_wide_no_dead[selected_columns].copy()


# =============================================================================
# TEMPORARY: REMOVE PLOTS
# =============================================================================

# As of 2025, the experiment changed so that only transects are monitored.
# Remove any rows where the transect begins with "P".
if EXCLUDE_PLOTS:
    coral_df_selected = coral_df_selected[
        ~coral_df_selected["transect"].str.startswith("P", na=False)
    ].copy()


# =============================================================================
# RENAME COLUMNS FOR THE DATA ENTRY SHEET
# =============================================================================

coral_df_selected = coral_df_selected.rename(
    columns={
        "site": "Site",
        "habitat": "Hab",
        "transect": "Tran",
        "taxa": "Taxa",
        "x": "X",
        "y": "Y",
        "z": "Z",
        length_col: f"L{PREVIOUS_SURVEY_YEAR % 100:02d}",
        width_col: f"W{PREVIOUS_SURVEY_YEAR % 100:02d}",
        height_col: f"H{PREVIOUS_SURVEY_YEAR % 100:02d}",
        notes_col: f"Notes{PREVIOUS_SURVEY_YEAR % 100:02d}",
    }
)

# Rename the 2-years-ago columns if INCLUDE_2_YEARS_AGO is True
if INCLUDE_2_YEARS_AGO:
    coral_df_selected = coral_df_selected.rename(
        columns={
            length_col_2_yrs_ago: f"L{TWO_YEARS_AGO % 100:02d}",
            width_col_2_yrs_ago: f"W{TWO_YEARS_AGO % 100:02d}",
            height_col_2_yrs_ago: f"H{TWO_YEARS_AGO % 100:02d}",
            notes_col_2_yrs_ago: f"Notes{TWO_YEARS_AGO % 100:02d}",
        }
    )


# =============================================================================
# CONVERT MEASUREMENT COLUMNS TO NUMERIC
# =============================================================================

numeric_columns = [
    "X",
    "Y",
    "Z",
    f"L{PREVIOUS_SURVEY_YEAR % 100:02d}",
    f"W{PREVIOUS_SURVEY_YEAR % 100:02d}",
    f"H{PREVIOUS_SURVEY_YEAR % 100:02d}",
]

for column in numeric_columns:
    coral_df_selected[column] = pd.to_numeric(
        coral_df_selected[column],
        errors="coerce"
    )


# =============================================================================
# ADD BLANK FIELDS FOR CURRENT SURVEY
# =============================================================================

current_year_short = CURRENT_SURVEY_YEAR % 100

coral_df_for_datasheet = coral_df_selected.copy()

coral_df_for_datasheet[f"L{current_year_short:02d}"] = ""
coral_df_for_datasheet[f"W{current_year_short:02d}"] = ""
coral_df_for_datasheet[f"H{current_year_short:02d}"] = ""
coral_df_for_datasheet[f"Notes{current_year_short:02d}"] = ""


# =============================================================================
# CREATE EXCEL WORKBOOK
# =============================================================================

wb = Workbook()

# Remove the default worksheet.
default_sheet = wb.active
wb.remove(default_sheet)


# =============================================================================
# CREATE ONE SHEET PER SITE / HABITAT / TRANSECT
# =============================================================================

unique_combos = (
    coral_df_for_datasheet[["Site", "Hab", "Tran"]]
    .drop_duplicates()
    .sort_values(["Site", "Hab", "Tran"])
)


for _, combo in unique_combos.iterrows():

    site = combo["Site"]
    hab = combo["Hab"]
    tran = combo["Tran"]

    sheet_name = f"{site}_{hab}_{tran}"

    # Excel worksheet names cannot exceed 31 characters.
    # Truncate if necessary.
    sheet_name = sheet_name[:31]

    # Filter to this Site/Hab/Tran combination.
    data_i = coral_df_for_datasheet[
        (coral_df_for_datasheet["Site"] == site)
        & (coral_df_for_datasheet["Hab"] == hab)
        & (coral_df_for_datasheet["Tran"] == tran)
    ].copy()

    # Remove Site, Hab, and Tran from the actual data table.
    data_i = data_i.drop(columns=["Site", "Hab", "Tran"])

    # Create worksheet.
    ws = wb.create_sheet(title=sheet_name)

    # -------------------------------------------------------------------------
    # HEADER
    # -------------------------------------------------------------------------

    header_text = (
        f"{sheet_name} | Date: ________________ | "
        f"Obs: __________________"
    )

    ws.cell(row=1, column=1, value=header_text)

    # Style header
    header_cell = ws.cell(row=1, column=1)
    header_cell.font = Font(size=12, bold=True)
    header_cell.alignment = Alignment(horizontal="left")

    # Merge header across all data columns.
    ws.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=len(data_i.columns)
    )

    # -------------------------------------------------------------------------
    # DATA
    # -------------------------------------------------------------------------

    # Write column headers in row 3.
    for column_index, column_name in enumerate(data_i.columns, start=1):
        cell = ws.cell(
            row=3,
            column=column_index,
            value=column_name
        )

        cell.font = Font(bold=True)

    # Write data beginning in row 4.
    for row_index, row in enumerate(
        data_i.itertuples(index=False, name=None),
        start=4
    ):
        for column_index, value in enumerate(row, start=1):

            # Convert values to strings to avoid Excel type issues,
            # matching the behavior of the original R code.
            if pd.isna(value):
                value = ""
            else:
                value = str(value)

            ws.cell(
                row=row_index,
                column=column_index,
                value=value
            )

    # -------------------------------------------------------------------------
    # COLUMN WIDTHS
    # -------------------------------------------------------------------------

    for column_index, column_name in enumerate(data_i.columns, start=1):

        column_letter = get_column_letter(column_index)

        # Basic automatic-ish width.
        max_length = len(str(column_name))

        for cell in ws[column_letter]:
            if cell.value is not None:
                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

        # Don't let columns become excessively wide.
        ws.column_dimensions[column_letter].width = min(
            max_length + 2,
            30
        )


# =============================================================================
# SAVE WORKBOOK
# =============================================================================

wb.save(OUTPUT_FILE)

print(f"Datasheet created: {OUTPUT_FILE}")
